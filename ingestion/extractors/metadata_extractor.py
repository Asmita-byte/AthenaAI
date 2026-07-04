from pathlib import Path

import fitz

from backend.core.logging import get_logger

logger = get_logger(__name__)


class MetadataExtractor:

    def extract(self, file_path: Path) -> dict:
        extension = file_path.suffix.lower()

        if extension == ".pdf":
            return self._extract_pdf_metadata(file_path)
        elif extension == ".docx":
            return self._extract_docx_metadata(file_path)
        elif extension == ".pptx":
            return self._extract_pptx_metadata(file_path)
        elif extension in (".xlsx", ".csv"):
            return self._extract_spreadsheet_metadata(file_path)
        else:
            return self._extract_basic_metadata(file_path)

    def _extract_pdf_metadata(self, file_path: Path) -> dict:
        try:
            with fitz.open(str(file_path)) as pdf:
                meta = pdf.metadata or {}
                return {
                    "title": meta.get("title", ""),
                    "author": meta.get("author", ""),
                    "subject": meta.get("subject", ""),
                    "creator": meta.get("creator", ""),
                    "total_pages": pdf.page_count,
                    "file_size_bytes": file_path.stat().st_size,
                    "extension": ".pdf",
                }
        except Exception as e:
            logger.warning("PDF metadata extraction failed", error=str(e))
            return self._extract_basic_metadata(file_path)

    def _extract_docx_metadata(self, file_path: Path) -> dict:
        try:
            from docx import Document
            doc = Document(str(file_path))
            props = doc.core_properties
            return {
                "title": props.title or "",
                "author": props.author or "",
                "subject": props.subject or "",
                "created": str(props.created or ""),
                "modified": str(props.modified or ""),
                "file_size_bytes": file_path.stat().st_size,
                "extension": ".docx",
            }
        except Exception as e:
            logger.warning("DOCX metadata extraction failed", error=str(e))
            return self._extract_basic_metadata(file_path)

    def _extract_pptx_metadata(self, file_path: Path) -> dict:
        try:
            from pptx import Presentation
            prs = Presentation(str(file_path))
            props = prs.core_properties
            return {
                "title": props.title or "",
                "author": props.author or "",
                "total_slides": len(prs.slides),
                "file_size_bytes": file_path.stat().st_size,
                "extension": ".pptx",
            }
        except Exception as e:
            logger.warning("PPTX metadata extraction failed", error=str(e))
            return self._extract_basic_metadata(file_path)

    def _extract_spreadsheet_metadata(self, file_path: Path) -> dict:
        try:
            import pandas as pd
            if file_path.suffix.lower() == ".csv":
                df = pd.read_csv(str(file_path), nrows=0)
                sheets = {"Sheet1": df}
            else:
                sheets = pd.read_excel(str(file_path), sheet_name=None, nrows=0)
            return {
                "total_sheets": len(sheets),
                "sheet_names": list(sheets.keys()),
                "file_size_bytes": file_path.stat().st_size,
                "extension": file_path.suffix.lower(),
            }
        except Exception as e:
            logger.warning("Spreadsheet metadata extraction failed", error=str(e))
            return self._extract_basic_metadata(file_path)

    def _extract_basic_metadata(self, file_path: Path) -> dict:
        return {
            "filename": file_path.name,
            "file_size_bytes": file_path.stat().st_size,
            "extension": file_path.suffix.lower(),
        }